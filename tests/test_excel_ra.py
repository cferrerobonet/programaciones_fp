"""Excel de RA (F9, ADR-041): el patrón y el diseño del Excel que acompaña a cada programación (la plantilla de
importación al cuaderno por RA de Educamos). Los textos son los del catálogo (oficiales); el resto, ficticio."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from conftest import programacion_minima
from programaciones import documento, excel_ra, exportacion


def libro(prog, cat):
    return load_workbook(BytesIO(excel_ra.generar(prog, cat)))


@pytest.fixture
def sea(cat):
    p = programacion_minima(cat, "SEA", 2, "0517")
    p["empresa"] = {"2": 12, "3": 13, "6": 0}   # el 6 va a empresa sin horas (modalidad Mixto o Empresa)
    return p


def test_cabecera_y_hojas(sea, cat):
    wb = libro(sea, cat)
    ws = wb.worksheets[0]
    assert wb.sheetnames == ["Procesos en instalaciones de in", "Leyenda"]      # 31 caracteres como mucho
    assert (ws["A1"].value, ws["C1"].value) == ("Módulo", "Procesos en instalaciones de infraestructuras comunes de "
                                                          "telecomunicaciones")
    assert (ws["A2"].value, ws["C2"].value) == ("Clase", "2º CFGS Sistemas Electrotécnicos y Automatizados")
    assert [ws.cell(5, c).value for c in (1, 2, 8, 9, 10, 11, 12)] == [
        "Reducido", "Nombre / Descripción ", "RA Asociado", "Visible", "Esencial", "RA Empresa", "Horas"]
    assert wb["Leyenda"]["B1"].value == "LEYENDA" and wb["Leyenda"]["B8"].value.startswith("Horas:")


def test_ra_y_criterios_como_en_el_excel_de_siempre(sea, cat):
    ws = libro(sea, cat).worksheets[0]
    ficha = cat.ficha("SEA", "0517", 2)
    n_ra, n_crit = len(ficha.ra), sum(len(r["criterios"]) for r in ficha.ra)
    assert ws.max_row == 5 + n_ra + n_crit
    fila = lambda r: [ws.cell(r, c).value for c in (1, 8, 9, 10, 11, 12)]   # noqa: E731
    assert ws["B6"].value.startswith("RA1- ") and fila(6) == ["RA1", None, "X", None, None, None]
    assert fila(7) == ["RA2", None, "X", None, "X", 12] and fila(8) == ["RA3", None, "X", None, "X", 13]
    assert fila(11) == ["RA6", None, "X", None, "X", None]                    # a empresa sin horas: sin número
    primero = 6 + n_ra
    assert fila(primero) == ["RA1.a", "RA1", "X", None, None, None] and ws.cell(primero, 2).value.startswith("a) ")
    # el texto, el del documento: con las erratas del BOE corregidas (ADR-018)
    assert "de acuerdo con la normativa" in ws.cell(primero, 2).value
    assert all(f"B{r}:G{r}" in {str(x) for x in ws.merged_cells.ranges} for r in range(6, ws.max_row + 1))
    assert isinstance(ws["L7"].value, int)                                     # Educamos exige que sea numérica


def test_la_exportacion_lo_escribe_con_el_nombre_de_la_programacion(sea, cat, curso, tmp_path):
    res = exportacion.exportar(sea, cat, curso, tmp_path, ["docx", "xlsx"], documento.Opciones())
    assert [p.name for p in res.archivos] == [
        "SEA 2 - 0517 - Procesos en instalaciones de infraestructuras comunes de telecomunicaciones.docx",
        "SEA 2 - 0517 - Procesos en instalaciones de infraestructuras comunes de telecomunicaciones.xlsx"]


def test_los_ambitos_no_tienen_excel_de_ra(cat, curso, tmp_path):
    from test_grado_basico import programacion_ambito

    amb = programacion_ambito(cat)
    assert not excel_ra.aplica(cat.ficha("EE", "3164", 2))
    with pytest.raises(excel_ra.SinExcel, match="ámbitos"):
        excel_ra.generar(amb, cat)
    res = exportacion.exportar(amb, cat, curso, tmp_path, ["docx", "xlsx"], documento.Opciones())
    assert [p.suffix for p in res.archivos] == [".docx"]


def test_el_formato_existe_pero_no_se_marca_de_serie(datos):
    from programaciones.ajustes import FORMATOS, GestorAjustes

    assert "xlsx" in FORMATOS and GestorAjustes(datos).valores.formatos == ["docx", "pdf"]
