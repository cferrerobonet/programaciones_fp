"""Excel de RA de una programación (F9, ADR-041): la plantilla de importación al cuaderno por RA de Educamos, con el
mismo patrón y diseño que el Excel que acompaña a cada programación en su carpeta `01 Programación` (CarlosFB,
2026-10-19: «no cambia nada, es idéntico patrón y diseño»).

Hoja del módulo (título: su nombre, 31 caracteres como mucho):
- Filas 1 y 2: «Módulo» y «Clase» («2º CFGS Sistemas Electrotécnicos y Automatizados»).
- Fila 5: Reducido · Nombre / Descripción (B:G combinadas) · RA Asociado · Visible · Esencial · RA Empresa · Horas.
- Desde la 6, primero los RA («RA1», «RA1- texto») y después sus criterios («RA1.a», «a) texto», RA asociado «RA1»).
  Visible: X en todos. Esencial: en blanco. RA Empresa: X en los RA que van a la empresa (tabla 4.1) y, si tienen
  horas, su número en Horas (numérico: Educamos rechaza la fila si no lo es).
Hoja «Leyenda»: qué es cada columna.

Los textos son los del documento (`Ficha.ra`): con las erratas del BOE ya corregidas (ADR-018). Los ámbitos de GB no
tienen RA sino competencias específicas: para ellos no hay Excel de RA (`aplica`).
"""
from __future__ import annotations

from io import BytesIO

from programaciones.catalogo import Catalogo, Ficha

NOMBRE_HOJA_MAX = 31   # límite de Excel
LEYENDA = [
    "LEYENDA",
    "Reducido: Campo alfanumérico para definir el reducido (RA1, RA1.a, etc.). Obligatorio. Máx 10 caracteres.",
    "Nombre /Descripción: Nombre del RA o Criterio. Obligatorio.",
    "RA asociado: Para criterios, rellenar con el código del RA padre (RA1, RA2...). Vacío = RA.",
    "Visible: X si debe ser visible para las familias.",
    "Esencial: X si el RA/Criterio es esencial (impide aprobar el módulo si no se supera).",
    "RA Empresa: X si el RA será evaluado en la empresa de prácticas.",
    "Horas: Número de horas asignadas al RA de Empresa (solo para RA Empresa).",
]
CABECERA = {1: "Reducido", 2: "Nombre / Descripción ", 8: "RA Asociado", 9: "Visible", 10: "Esencial",
            11: "RA Empresa", 12: "Horas"}
ANCHOS = {"A": 17.3, "B": 4.4, "C": 8.4, "D": 8.4, "E": 8.4, "F": 8.4, "G": 45.6, "H": 12.4, "I": 8.0, "J": 8.0,
          "K": 11.0, "L": 7.0}
TITULOS = ("Técnico Superior en ", "Técnico Básico en ", "Técnico en ")


class SinExcel(ValueError):
    """El módulo no tiene Excel de RA (un ámbito de GB) o el catálogo no tiene sus RA."""


def aplica(ficha: Ficha) -> bool:
    return not ficha.es_ambito and bool(ficha.ra)


def clase(ficha: Ficha) -> str:
    """«2º CFGS Sistemas Electrotécnicos y Automatizados»: curso, grado y el título sin «Técnico Superior en»."""
    nombre = ficha.titulo
    for prefijo in TITULOS:
        if nombre.startswith(prefijo):
            nombre = nombre[len(prefijo):]
            break
    return f"{ficha.curso}º CF{ficha.grado} {nombre}"


def filas(prog: dict, ficha: Ficha) -> list[tuple]:
    """(reducido, nombre, RA asociado, visible, esencial, RA empresa, horas) de cada fila, en orden."""
    empresa = {int(k): int(v or 0) for k, v in (prog.get("empresa") or {}).items()}
    out = [(f"RA{r['n']}", f"RA{r['n']}- {r['texto']}", None, "X", None, "X" if r["n"] in empresa else None,
            empresa.get(r["n"]) or None) for r in ficha.ra]
    for r in ficha.ra:
        out += [(f"RA{r['n']}.{c['letra']}", f"{c['letra']}) {c['texto']}", f"RA{r['n']}", "X", None, None, None)
                for c in r["criterios"]]
    return out


def generar(prog: dict, cat: Catalogo) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    ficha = cat.ficha(prog["ciclo"], prog["modulo"], prog.get("curso"))
    if not aplica(ficha):
        raise SinExcel("Los ámbitos de Grado Básico no tienen Excel de RA" if ficha.es_ambito
                       else "El catálogo no tiene los RA de este módulo")
    wb = Workbook()
    ws = wb.active
    ws.title = ficha.nombre[:NOMBRE_HOJA_MAX]
    for fila, (etiqueta, valor) in enumerate((("Módulo", ficha.nombre), ("Clase", clase(ficha))), start=1):
        ws.cell(fila, 1, etiqueta).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(fila, 3, valor).font = Font(name="Calibri", size=11)
    for col, texto in CABECERA.items():
        c = ws.cell(5, col, texto)
        c.font = Font(name="Calibri", size=11, bold=True)
        c.alignment = Alignment(horizontal="left" if col == 2 else "center")
    for i, valores in enumerate(filas(prog, ficha), start=6):
        for col, valor in zip((1, 2, 8, 9, 10, 11, 12), valores, strict=True):
            c = ws.cell(i, col, valor)
            c.font = Font(name="Calibri", size=11)
            if col == 2:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
    for col, ancho in ANCHOS.items():
        ws.column_dimensions[col].width = ancho

    ley = wb.create_sheet("Leyenda")
    for i, texto in enumerate(LEYENDA, start=1):
        ley.cell(i, 2, texto).font = Font(name="Calibri", size=11)
    ley.column_dimensions["B"].width = 90.0

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
